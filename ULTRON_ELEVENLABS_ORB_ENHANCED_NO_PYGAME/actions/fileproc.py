"""Process a file: describe/OCR/summarize/transcribe (AI) or resize/convert/extract (local)."""
import json
import shutil
import subprocess
from pathlib import Path


def _out(src: Path, suffix: str, ext=None) -> Path:
    return src.parent / f"{src.stem}_{suffix}{ext or src.suffix}"


def _detect(p: Path) -> str:
    e = p.suffix.lower().lstrip(".")
    for kind, exts in {
        "image": {"jpg", "jpeg", "png", "gif", "webp", "bmp", "tiff", "heic"},
        "video": {"mp4", "avi", "mov", "mkv", "webm", "m4v"},
        "audio": {"mp3", "wav", "ogg", "m4a", "aac", "flac"},
        "code": {"py", "js", "ts", "html", "css", "java", "c", "cpp", "go", "rs", "sh", "sql"},
        "archive": {"zip", "tar", "gz", "7z", "rar"},
        "data": {"csv", "tsv", "xlsx", "xls"},
    }.items():
        if e in exts:
            return kind
    return {"pdf": "pdf", "docx": "docx", "doc": "docx", "txt": "text", "md": "text",
            "json": "json", "pptx": "pptx"}.get(e, "unknown")


def _ai(prompt):
    from ai_helper import complete_text
    return complete_text(prompt)


def file_process(file_path: str, action: str = "", instruction: str = "",
                 format: str = "", width: int = 0, height: int = 0,
                 quality: int = 70, start: str = "", end: str = "") -> str:
    p = Path(file_path).expanduser()
    if not p.is_file():
        return f"File not found: {file_path}"
    kind = _detect(p)
    a = (action or "").lower().strip()

    try:
        if kind == "image":
            from PIL import Image
            if a in ("", "describe", "ocr", "analyze"):
                import io
                img = Image.open(p).convert("RGB")
                img.thumbnail((1280, 1280))
                buf = io.BytesIO(); img.save(buf, "JPEG", quality=85)
                from ai_helper import analyze_image
                q = instruction or {"ocr": "Extract all text from this image.",
                                    "analyze": "Analyze this image in detail."}.get(a, "Describe this image.")
                return analyze_image(buf.getvalue(), q, "image/jpeg")
            if a == "resize":
                img = Image.open(p); w, h = img.size
                size = (width, height) if width and height else \
                       (width, int(h * width / w)) if width else \
                       (int(w * height / h), height) if height else None
                if not size:
                    return "Give width and/or height."
                o = _out(p, f"{size[0]}x{size[1]}")
                img.resize(size, Image.LANCZOS).save(o)
                return f"Resized to {size[0]}x{size[1]}: {o.name}"
            if a == "convert":
                fmt = (format or "png").lower()
                o = _out(p, "conv", f".{fmt}")
                im = Image.open(p)
                (im.convert("RGB") if fmt in ("jpg", "jpeg") else im).save(o)
                return f"Converted to {fmt}: {o.name}"
            if a == "compress":
                o = _out(p, "compressed", ".jpg")
                Image.open(p).convert("RGB").save(o, "JPEG", quality=quality, optimize=True)
                return f"Compressed: {o.name}"
            return "Image actions: describe, ocr, analyze, resize, convert, compress."

        if kind == "pdf":
            text = ""
            try:
                import pdfplumber
                with pdfplumber.open(p) as pdf:
                    text = "\n".join((pg.extract_text() or "") for pg in pdf.pages)
            except ImportError:
                try:
                    import PyPDF2
                    with open(p, "rb") as f:
                        text = "\n".join(pg.extract_text() or "" for pg in PyPDF2.PdfReader(f).pages)
                except ImportError:
                    return "Install pdfplumber or PyPDF2 to read PDFs."
            if not text.strip():
                return "No extractable text (may be a scanned PDF)."
            if a == "extract_text":
                o = _out(p, "text", ".txt"); o.write_text(text, encoding="utf-8")
                return f"Extracted {len(text)} chars to {o.name}"
            return _ai(f"{instruction or 'Summarize this document'}:\n\n{text[:40000]}")

        if kind in ("text", "docx"):
            if kind == "docx":
                from docx import Document
                text = "\n".join(par.text for par in Document(p).paragraphs)
            else:
                text = p.read_text(encoding="utf-8", errors="ignore")
            if a == "word_count":
                return f"{len(text.split())} words, {len(text)} chars, {text.count(chr(10))+1} lines."
            return _ai(f"{instruction or 'Summarize this'}:\n\n{text[:40000]}")

        if kind == "data":
            import pandas as pd
            df = pd.read_csv(p) if p.suffix.lower() in (".csv", ".tsv") else pd.read_excel(p)
            if a == "stats":
                return df.describe(include="all").to_string()[:3000]
            if a == "info":
                return f"{len(df)} rows, {len(df.columns)} cols: {', '.join(map(str, df.columns))}"
            if a == "convert":
                fmt = (format or "csv").lower()
                o = _out(p, "conv", f".{fmt}")
                {"csv": lambda: df.to_csv(o, index=False),
                 "xlsx": lambda: df.to_excel(o, index=False),
                 "json": lambda: df.to_json(o, orient="records", indent=2)}.get(fmt, lambda: df.to_csv(o, index=False))()
                return f"Converted to {fmt}: {o.name}"
            return _ai(f"Analyze this dataset ({len(df)} rows, cols {list(df.columns)}):\n"
                       f"{df.head(40).to_string()}\n\n{instruction or 'Give key insights.'}")

        if kind == "json":
            data = json.loads(p.read_text(encoding="utf-8"))
            if a == "validate":
                return f"Valid JSON ({type(data).__name__})."
            if a == "format":
                o = _out(p, "formatted", ".json")
                o.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
                return f"Formatted: {o.name}"
            return _ai(f"{instruction or 'Analyze this JSON'}:\n{json.dumps(data, indent=2)[:8000]}")

        if kind == "code":
            content = p.read_text(encoding="utf-8", errors="ignore")[:30000]
            task = {"review": "Review this code for bugs and improvements",
                    "explain": "Explain what this code does",
                    "fix": "Fix bugs and return the corrected code",
                    "document": "Add documentation/comments",
                    "test": "Write unit tests for this code"}.get(a, instruction or "Explain this code")
            return _ai(f"{task}:\n\n```\n{content}\n```")

        if kind == "audio":
            if a in ("", "transcribe"):
                return _transcribe_audio(p)
            return _ffmpeg_media(p, a, format, start, end, quality)

        if kind == "video":
            return _ffmpeg_media(p, a or "info", format, start, end, quality)

        if kind == "archive":
            if a == "extract":
                dest = p.parent / p.stem
                dest.mkdir(exist_ok=True)
                shutil.unpack_archive(str(p), str(dest))
                return f"Extracted to {dest}"
            import zipfile
            if p.suffix.lower() == ".zip":
                with zipfile.ZipFile(p) as z:
                    names = z.namelist()
                return f"{len(names)} entries:\n" + "\n".join(names[:40])
            return "For non-zip archives use action=extract."

        if kind == "pptx":
            from pptx import Presentation
            text = "\n".join(sh.text for sl in Presentation(p).slides
                             for sh in sl.shapes if hasattr(sh, "text"))
            return _ai(f"{instruction or 'Summarize this presentation'}:\n{text[:30000]}")

        return f"Unsupported file type: {p.suffix}"
    except ImportError as e:
        return f"A library is missing for that: {e}"
    except Exception as e:
        return f"Processing failed: {e}"


def _transcribe_audio(p: Path) -> str:
    try:
        from faster_whisper import WhisperModel
        model = WhisperModel("base", device="cpu", compute_type="int8")
        segs, _ = model.transcribe(str(p))
        return " ".join(s.text for s in segs).strip() or "(no speech detected)"
    except ImportError:
        return "Install faster-whisper to transcribe audio: pip install faster-whisper"
    except Exception as e:
        return f"Transcription failed: {e}"


def _ffmpeg_media(p: Path, action: str, fmt: str, start: str, end: str, quality: int) -> str:
    if not shutil.which("ffmpeg"):
        return "ffmpeg isn't installed; needed for audio/video processing."
    a = action.lower()
    try:
        if a == "info":
            r = subprocess.run(["ffprobe", "-v", "quiet", "-show_format", "-show_streams",
                                "-print_format", "json", str(p)], capture_output=True, text=True, timeout=20)
            d = json.loads(r.stdout)
            dur = float(d.get("format", {}).get("duration", 0))
            return f"{int(dur//60)}m {int(dur%60)}s, {p.stat().st_size/1024**2:.1f} MB"
        if a == "extract_audio":
            o = _out(p, "audio", ".mp3")
            subprocess.run(["ffmpeg", "-i", str(p), "-q:a", "0", "-map", "a", str(o), "-y"],
                           capture_output=True, timeout=300)
            return f"Audio extracted: {o.name}"
        if a == "trim":
            o = _out(p, "trim")
            cmd = ["ffmpeg", "-i", str(p), "-ss", start or "0"]
            if end:
                cmd += ["-to", end]
            cmd += ["-c", "copy", str(o), "-y"]
            subprocess.run(cmd, capture_output=True, timeout=600)
            return f"Trimmed: {o.name}"
        if a == "convert":
            o = _out(p, "conv", f".{fmt or 'mp4'}")
            subprocess.run(["ffmpeg", "-i", str(p), str(o), "-y"], capture_output=True, timeout=1800)
            return f"Converted: {o.name}"
        return "Media actions: info, extract_audio, trim, convert."
    except Exception as e:
        return f"ffmpeg failed: {e}"
